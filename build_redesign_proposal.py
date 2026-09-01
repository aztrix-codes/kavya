from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "Kavya_Portfolio_Redesign_Proposal.docx"


PALETTES = [
    {
        "name": "A. Scientific Vibrance",
        "position": "Recommended primary direction",
        "colors": [
            ("Deep Ocean", "#073B4C", "Header, footer, dark text"),
            ("Aqua Signal", "#00B4D8", "Primary CTA, links, hover states"),
            ("Fresh Teal", "#2EC4B6", "Data highlights and badges"),
            ("Warm Coral", "#FF6B6B", "Selective emphasis and alerts"),
            ("Clean Mist", "#F6FBFC", "Page background"),
        ],
        "why": "Best match for groundwater, Earth observation, and AI. It feels modern, active, and scientific without becoming playful.",
    },
    {
        "name": "B. Earth Observation",
        "position": "Grounded academic alternative",
        "colors": [
            ("Forest", "#1B4332", "Navigation, section anchors"),
            ("Satellite Green", "#52B788", "Buttons and active states"),
            ("River Blue", "#168AAD", "Links and charts"),
            ("Terrain Gold", "#F4A261", "Warm contrast accents"),
            ("Paper", "#FAFAF7", "Background"),
        ],
        "why": "Connects directly to environmental science, water security, maps, and policy. It is vibrant but still mature.",
    },
    {
        "name": "C. Atlas Dark",
        "position": "Premium dark-mode led concept",
        "colors": [
            ("Night Map", "#0B1020", "Dark background"),
            ("Electric Cyan", "#22D3EE", "Primary highlights"),
            ("Lime Data", "#A3E635", "Metrics and success states"),
            ("Violet Model", "#8B5CF6", "AI/ML accents"),
            ("Cloud Text", "#E5E7EB", "Text on dark"),
        ],
        "why": "Good if she wants a high-impact, technical portfolio with map/dashboard energy. Needs careful contrast and restraint.",
    },
]


STYLE_ROUTES = [
    (
        "Scientific Vibrance",
        "Recommended",
        "Clean academic layout with bright data accents, map textures, water/terrain gradients, and polished cards.",
        "Best balance: professional enough for professors and collaborators, vibrant enough to satisfy the redesign request.",
    ),
    (
        "Minimal Research Editorial",
        "Safe option",
        "More white space, strong typography, large research visuals, and restrained color. Similar to a journal profile.",
        "Useful if the client worries that too much color may reduce academic seriousness.",
    ),
    (
        "Material UI Dashboard",
        "Functional option",
        "Tabbed research areas, cards, filters, metrics, publication lists, and admin-friendly components.",
        "Strong fit for a Next.js build where she can edit sections, add projects, and update tabs.",
    ),
    (
        "Neo-Brutal Accent",
        "Use carefully",
        "Bold outlines, strong contrast, high-energy color blocks, and oversized labels.",
        "Can look memorable, but too much will feel less academic. Use only for selected banners or project cards.",
    ),
    (
        "3D / Interactive Atlas",
        "Premium add-on",
        "Subtle globe/map layers, animated terrain cards, risk-map overlays, and interactive data visuals.",
        "Good for a premium phase, but should not delay the core portfolio rebuild.",
    ),
]


def set_cell_fill(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color.replace("#", ""))
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def set_cell_border(cell, color="D9E2E7", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            first, rest = item
            r = p.add_run(first)
            r.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        if isinstance(item, tuple):
            first, rest = item
            r = p.add_run(first)
            r.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_callout(doc, label, text, fill="EAF8FB", border="00B4D8"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.3)
    cell = table.cell(0, 0)
    cell.width = Inches(6.3)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_fill(cell, fill)
    set_cell_border(cell, border, "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("073B4C")
    p.add_run(text)
    doc.add_paragraph()


def add_palette(doc, palette):
    doc.add_heading(palette["name"], level=2)
    add_para(doc, palette["position"], style="Intense Quote")
    add_para(doc, palette["why"])
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.25), Inches(1.25), Inches(1.25), Inches(1.25), Inches(1.25)]
    for i, width in enumerate(widths):
        table.columns[i].width = width
    for idx, (name, hex_color, use) in enumerate(palette["colors"]):
        cell = table.cell(0, idx)
        cell.width = widths[idx]
        set_cell_fill(cell, hex_color)
        luminance_dark = hex_color.lower() in ["#073b4c", "#1b4332", "#0b1020", "#168aad", "#8b5cf6"]
        set_cell_text(cell, "{}\n{}\n{}".format(name, hex_color, use), bold=True, color="FFFFFF" if luminance_dark else "111827", size=8)
        set_cell_border(cell, "FFFFFF", "8")
    doc.add_paragraph()


def add_two_col_table(doc, rows, widths=(1.75, 4.55), header=None):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(widths[0])
    table.columns[1].width = Inches(widths[1])
    if header:
        cells = table.add_row().cells
        for i, text in enumerate(header):
            cells[i].width = Inches(widths[i])
            set_cell_fill(cells[i], "073B4C")
            set_cell_text(cells[i], text, bold=True, color="FFFFFF", size=9)
            set_cell_border(cells[i], "073B4C", "6")
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        set_cell_fill(cells[0], "F2FBFC")
        set_cell_text(cells[0], left, bold=True, color="073B4C", size=9)
        set_cell_text(cells[1], right, size=9)
        set_cell_border(cells[0])
        set_cell_border(cells[1])
    doc.add_paragraph()


def apply_document_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string("1F2937")
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in [
        ("Title", 24, "073B4C", 0, 6),
        ("Subtitle", 11, "475569", 0, 12),
        ("Heading 1", 16, "073B4C", 14, 7),
        ("Heading 2", 13, "00A6C8", 10, 5),
        ("Heading 3", 11.5, "1B4332", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for list_style in ["List Bullet", "List Number"]:
        style = styles[list_style]
        style.font.name = "Calibri"
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.15


def build():
    doc = Document()
    apply_document_styles(doc)

    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("Kavya Agrawal Portfolio Redesign Brief")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("64748B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Prepared for discussion - Next.js portfolio redesign")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string("64748B")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Kavya Agrawal Portfolio Redesign Proposal")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Creative direction, color systems, graphics plan, and editable Next.js rebuild scope")

    add_callout(
        doc,
        "Recommended direction",
        "Build a vibrant but credible scientific portfolio: water, satellite, terrain, and data-visual accents layered over a clean academic structure. The safest pitch is not 'make it colorful'; it is 'make it more memorable while preserving research trust.'",
        "EAF8FB",
        "00B4D8",
    )

    doc.add_heading("1. Current Portfolio Reading", level=1)
    add_para(
        doc,
        "The existing portfolio is a static HTML academic site for Kavya Agrawal, a Ph.D. scholar at TERI School of Advanced Studies. It is clean and professional, with light/dark themes, restrained cards, blue links, and sections for About, Education & Experience, Research, Projects, Publications, Research Notes, CV, and Contact.",
    )
    add_para(
        doc,
        "The content already gives strong design cues: groundwater contamination, geospatial AI, Earth observation, uranium/fluoride/arsenic risk mapping, population exposure, governance diagnostics, and cryosphere/hydrology research.",
    )
    add_bullets(
        doc,
        [
            ("Main issue: ", "The current visual language is too plain for a client asking for vibrant color, graphics, and a more memorable impression."),
            ("Main opportunity: ", "Use her scientific subject matter as the visual system: maps, water layers, satellite grids, risk gradients, and research data cards."),
            ("Important constraint: ", "Because she is an academic researcher, the redesign should stay credible and readable. Avoid a generic colorful landing page that feels unrelated to her work."),
        ]
    )

    doc.add_heading("2. Proposed Positioning", level=1)
    add_callout(
        doc,
        "Pitch line",
        "A modern research portfolio for a geospatial AI and environmental risk scholar - visually vibrant, data-rich, editable, and suitable for collaborations, fellowships, postdoctoral opportunities, and publication visibility.",
        "F2FFF8",
        "2EC4B6",
    )
    add_two_col_table(
        doc,
        [
            ("Visual personality", "Scientific, clear, warm, research-led, data-driven, trustworthy, and visually energetic."),
            ("Design references", "Research atlas, modern climate-tech site, clean Material UI dashboard, academic editorial profile, and selective neo-brutal accents."),
            ("Do more of", "Map-led graphics, section color coding, project filters, publication cards, research metrics, tabbed research themes, and interactive figure galleries."),
            ("Avoid", "Overly playful colors, heavy decorative animations, unrelated abstract blobs, dense text walls, and styling that makes the site look less academic."),
        ],
        header=("Area", "Recommendation"),
    )

    doc.add_heading("3. Color Scheme Samples", level=1)
    add_para(
        doc,
        "These palettes are designed to match Kavya's work rather than random bright colors. The first option is the recommended system because it connects water, risk mapping, and AI without losing academic seriousness.",
    )
    for palette in PALETTES:
        add_palette(doc, palette)

    doc.add_heading("4. Style Direction Options", level=1)
    rows = [(name + " (" + priority + ")", description + " " + fit) for name, priority, description, fit in STYLE_ROUTES]
    add_two_col_table(doc, rows, widths=(2.05, 4.25), header=("Style", "How to explain it"))

    doc.add_heading("5. Recommended Page Experience", level=1)
    add_numbered(
        doc,
        [
            ("Hero: ", "Use a stronger first screen with Kavya's name, role, research focus, profile photo, and a subtle map/water/data background. Add CTAs: View Research, Download CV, Contact."),
            ("Research themes: ", "Use colorful tabs for Groundwater, Geospatial AI, Exposure Assessment, Policy, and Cryosphere. Each tab can include summary, methods, datasets, outputs, and key visuals."),
            ("Projects: ", "Convert long project blocks into filterable cards with status labels, contaminant tags, map previews, methods, outcomes, and links to details."),
            ("Publications: ", "Create a polished publication archive with filters for Published, Under Review, Conference, Technical Reports, and featured work."),
            ("Graphics gallery: ", "Add a map and figure gallery using existing project images, with captions and research context."),
            ("CV and profile: ", "Keep academic credibility with a structured CV page, research portals, awards, workshops, teaching, and collaboration interests."),
        ]
    )

    doc.add_heading("6. Graphics and Visual Assets to Add", level=1)
    add_bullets(
        doc,
        [
            ("Map textures: ", "Faint India outline, terrain contours, water-network lines, satellite tile overlays, and risk heat-map gradients."),
            ("Data graphics: ", "Small metric cards for publications, research themes, datasets, study regions, and tools used."),
            ("Research icons: ", "Water droplet for groundwater, satellite for Earth observation, neural network for AI, shield/health for exposure risk, mountain for cryosphere."),
            ("Project previews: ", "Use existing project images as visual anchors instead of generic stock graphics."),
            ("Micro-interactions: ", "Soft card hover states, animated tab transitions, map zoom reveals, and counter animations for metrics."),
            ("Optional premium visual: ", "A lightweight interactive research atlas showing study themes and map thumbnails. Keep it fast and accessible."),
        ]
    )

    doc.add_heading("7. Next.js Rebuild and Editable Content Plan", level=1)
    add_para(
        doc,
        "The strongest technical proposal is to rebuild the site in Next.js so the portfolio becomes easier to maintain, faster to extend, and more professional. The client specifically wants the ability to add tabs and edit content; that should be handled through structured content, not by editing raw HTML.",
    )
    add_two_col_table(
        doc,
        [
            ("Frontend", "Next.js App Router, TypeScript, responsive layout, SEO metadata, image optimization, accessibility checks, and optional PWA support."),
            ("Styling", "Tailwind CSS with a controlled design-token system, or Material UI if she prefers a dashboard-like interface. Framer Motion can be added for subtle transitions."),
            ("Editable content", "Option 1: simple JSON/MDX content files for low-cost maintenance. Option 2: a headless CMS such as Sanity or Contentful for non-technical editing. Option 3: custom admin dashboard for the highest control."),
            ("Tabs and sections", "Schema-driven tabs where she can add Research Themes, Project Categories, Publications, Research Notes, Skills, or News without changing layout code."),
            ("Content model", "Profile, Research Theme, Project, Publication, News Item, Gallery Item, Skill/Tool, CV Entry, and Contact/Social Link."),
            ("Deployment", "Vercel or similar hosting with preview links, production deployment, analytics, and easy future updates."),
        ],
        header=("Layer", "Recommendation"),
    )

    doc.add_heading("8. What Can Be Offered to the Client", level=1)
    add_two_col_table(
        doc,
        [
            ("Base redesign", "New visual design, color system, responsive pages, improved navigation, cleaned content hierarchy, and updated project/publication presentation."),
            ("Next.js build", "Modern implementation, reusable components, SEO setup, image optimization, and deployment-ready code."),
            ("Editable system", "CMS or structured content setup so Kavya can add or edit tabs, projects, publications, and research notes."),
            ("Graphics package", "Map-themed backgrounds, scientific icon set, project thumbnails treatment, section illustrations, and consistent data badges."),
            ("Premium add-ons", "Interactive research atlas, filtered publications, downloadable CV styling, admin dashboard, animations, analytics, and blog/research-note editor."),
        ],
        header=("Offer", "Explanation for negotiation"),
    )

    doc.add_heading("9. Suggested Delivery Phases", level=1)
    add_two_col_table(
        doc,
        [
            ("Phase 1: Direction", "Approve palette, style route, homepage wireframe, and content structure. Output: visual direction and page map."),
            ("Phase 2: Design", "Design homepage, research/projects pages, publication cards, tabs, and mobile views. Output: design mockups or coded prototype."),
            ("Phase 3: Build", "Develop Next.js site, component system, responsive pages, SEO, and initial content migration."),
            ("Phase 4: Editing Setup", "Add CMS/content files/admin flow and train client on adding tabs, projects, publications, and notes."),
            ("Phase 5: Polish", "Performance, accessibility, metadata, deployment, final QA, and handover documentation."),
        ],
        header=("Phase", "Scope"),
    )

    doc.add_heading("10. Partner Talking Points", level=1)
    add_bullets(
        doc,
        [
            "The goal is not only a prettier portfolio. The goal is a more credible academic identity that is easier to update.",
            "The design will use her own research themes as the visual language: water, maps, satellite data, AI, contaminants, and climate risk.",
            "Vibrant color will be used intentionally: for section coding, data highlights, buttons, maps, and project statuses.",
            "Next.js gives a better foundation than static HTML for reusable tabs, editable sections, SEO, and future expansion.",
            "We can offer different levels depending on budget: visual refresh only, full Next.js rebuild, or rebuild with CMS/admin editing.",
        ]
    )

    doc.add_heading("11. Questions to Confirm Before Pricing", level=1)
    add_numbered(
        doc,
        [
            "Does Kavya want the site to feel more academic, more creative, or more dashboard-like?",
            "Should she edit content herself through a CMS/admin panel, or is developer-managed editing acceptable?",
            "Which content must be editable: tabs, projects, publications, research notes, CV, gallery images, or all of them?",
            "Does she want interactive maps/graphics now, or should that be a later premium phase?",
            "Will she provide updated photos, CV PDF, publication list, and final research images?",
            "What is the preferred hosting and domain setup?",
        ]
    )

    doc.add_heading("12. Recommended Final Proposal Package", level=1)
    add_callout(
        doc,
        "Best package to pitch",
        "Full Next.js redesign with Scientific Vibrance palette, tabbed research sections, filterable project/publication cards, map-led graphics, and an editable content system. This gives the client the vibrant feel she requested while solving the long-term maintenance problem.",
        "FFF7ED",
        "F4A261",
    )
    add_para(
        doc,
        "Suggested wording for the partner: We can redesign the portfolio around your research identity, not just add random colors. The new version can feel vibrant and modern while staying suitable for academic collaborators, fellowships, and postdoctoral opportunities. We can also rebuild it in Next.js so adding new tabs, publications, projects, and research notes becomes structured and easy.",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
