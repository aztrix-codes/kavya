from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "Kavya_Portfolio_Redesign_Proposal_Updated.docx"

PALETTES = [
    {
        "name": "A. Dreamy Aqua Bloom",
        "position": "Dreamy, soft, modern, and beautiful",
        "colors": [
            ("Deep Lagoon", "#102A43", "Headers and navigation"),
            ("Aurora Aqua", "#2DD4BF", "Primary buttons and links"),
            ("Soft Lavender", "#A78BFA", "Section accents and tabs"),
            ("Peach Glow", "#FB7185", "Highlights and featured cards"),
            ("Cloud White", "#F8FAFC", "Background"),
        ],
        "why": "Best if she wants the portfolio to feel elegant, fresh, and memorable. It gives a soft premium look while still matching water, research, and environmental themes.",
    },
    {
        "name": "B. Wow Satellite Glow",
        "position": "High-impact, visual, and energetic",
        "colors": [
            ("Deep Space", "#07111F", "Dark hero and footer"),
            ("Electric Cyan", "#06B6D4", "Main action color"),
            ("Vivid Lime", "#84CC16", "Metrics and data points"),
            ("Solar Amber", "#F59E0B", "Map/risk highlights"),
            ("Magenta Pulse", "#E879F9", "Rare wow accents"),
        ],
        "why": "Best for a dramatic homepage, glowing map graphics, and a strong first impression. Use carefully so the site feels premium, not flashy.",
    },
    {
        "name": "C. Professional Terra Blue",
        "position": "Most professional and safest for academia",
        "colors": [
            ("Academic Navy", "#0F172A", "Text, header, footer"),
            ("Research Blue", "#2563EB", "Links and CTAs"),
            ("Clean Emerald", "#10B981", "Success/data accents"),
            ("Terrain Gold", "#D97706", "Warm section highlights"),
            ("Porcelain", "#F8FAFC", "Background"),
        ],
        "why": "Best if the client wants color but does not want to risk losing academic credibility. It is polished, clean, and easy to maintain.",
    },
    {
        "name": "D. Dreamy Earth Pastel",
        "position": "Soft, graceful, calm, and personal",
        "colors": [
            ("Ink Grey", "#1F2937", "Main text"),
            ("Glacier Blue", "#7DD3FC", "Water and map sections"),
            ("Sage Green", "#86EFAC", "Environment accents"),
            ("Orchid Mist", "#C4B5FD", "AI/research tabs"),
            ("Blush Tint", "#FBCFE8", "Soft feature backgrounds"),
        ],
        "why": "Good if she wants a graceful personal portfolio with dreamy colors. It is less corporate and more personal, but still clean.",
    },
    {
        "name": "E. Premium Editorial Contrast",
        "position": "Bold, mature, and designer-looking",
        "colors": [
            ("Charcoal", "#111827", "Typography and dark blocks"),
            ("Cerulean", "#0EA5E9", "Research links and buttons"),
            ("Rose Red", "#E11D48", "Featured highlights"),
            ("Olive Signal", "#65A30D", "Environmental tags"),
            ("Warm Ivory", "#FFF7ED", "Editorial background"),
        ],
        "why": "Good for a more publication/editorial feel. It looks confident and premium, especially with large project images and strong typography.",
    },
]

def set_cell_fill(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color.replace("#", ""))
    tc_pr.append(shd)

def set_cell_border(cell, color="D9E2E7", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)

def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))

def add_para(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            r = p.add_run(item[0])
            r.bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)

def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        if isinstance(item, tuple):
            r = p.add_run(item[0])
            r.bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)

def add_callout(doc, label, text, fill, border):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.3)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_fill(cell, fill)
    set_cell_border(cell, border, "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("102A43")
    p.add_run(text)
    doc.add_paragraph()

def add_two_col_table(doc, rows, widths=(2.0, 4.3), header=("Area", "Recommendation")):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if header:
        cells = table.add_row().cells
        for i, text in enumerate(header):
            cells[i].width = Inches(widths[i])
            set_cell_fill(cells[i], "102A43")
            set_cell_text(cells[i], text, bold=True, color="FFFFFF", size=9)
            set_cell_border(cells[i], "102A43")
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        set_cell_fill(cells[0], "F1FBFC")
        set_cell_text(cells[0], left, bold=True, color="102A43", size=9)
        set_cell_text(cells[1], right, size=9)
        set_cell_border(cells[0])
        set_cell_border(cells[1])
    doc.add_paragraph()

def add_palette(doc, palette):
    doc.add_heading(palette["name"], level=2)
    add_para(doc, palette["position"], style="Intense Quote")
    add_para(doc, palette["why"])
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    darks = {"#102A43", "#07111F", "#0F172A", "#1F2937", "#111827", "#2563EB", "#E11D48", "#65A30D", "#0EA5E9"}
    for idx, (name, hex_color, use) in enumerate(palette["colors"]):
        cell = table.cell(0, idx)
        cell.width = Inches(1.25)
        set_cell_fill(cell, hex_color)
        text_color = "FFFFFF" if hex_color.upper() in darks else "111827"
        set_cell_text(cell, f"{name}\n{hex_color}\n{use}", bold=True, color=text_color, size=8)
        set_cell_border(cell, "FFFFFF", "8")
    doc.add_paragraph()

def apply_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string("1F2937")
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in [
        ("Title", 24, "102A43", 0, 6),
        ("Subtitle", 11, "475569", 0, 12),
        ("Heading 1", 16, "102A43", 14, 7),
        ("Heading 2", 13, "06B6D4", 10, 5),
        ("Heading 3", 11.5, "1B4332", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


doc = Document()
apply_styles(doc)
header = doc.sections[0].header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("Kavya Agrawal Portfolio Redesign Brief")
hr.font.size = Pt(8)
hr.font.color.rgb = RGBColor.from_string("64748B")
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("Updated color choices + Next.js editable portfolio plan")
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor.from_string("64748B")

title = doc.add_paragraph(style="Title")
title.add_run("Kavya Agrawal Portfolio Redesign Proposal")
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.add_run("Updated creative direction with dreamy, wow, and professional color scheme options")

add_callout(doc, "Recommended pitch", "Do not force one color scheme upfront. Present 3-5 polished visual directions and let Kavya choose the mood: dreamy, wow, professional, soft pastel, or premium editorial. The Next.js rebuild can also support a theme selector later.", "EAF8FB", "06B6D4")

doc.add_heading("1. Current Portfolio Reading", level=1)
add_para(doc, "The existing portfolio is a static HTML academic site for Kavya Agrawal, a Ph.D. scholar at TERI School of Advanced Studies. It is clean, professional, and content-rich, but the current visual language is restrained. The client request is for more vibrant colors, stronger graphics, and a design that feels more personal and memorable.")
add_bullets(doc, [
    ("Main opportunity: ", "Use Kavya's actual research world as the design language: water, satellite maps, terrain, groundwater risk, contaminants, AI models, and environmental health."),
    ("Important balance: ", "The portfolio can become vibrant without becoming childish or unacademic. Color should support research identity, not distract from it."),
    ("Recommended method: ", "Offer color schemes as mood options first, then choose one default theme for the final build."),
])

doc.add_heading("2. Color Scheme Choice Model", level=1)
add_para(doc, "Use this table in the negotiation discussion. It makes the color decision simple and gives the client control without making the project vague.")
add_two_col_table(doc, [
    ("Dreamy / graceful", "Choose Dreamy Aqua Bloom or Dreamy Earth Pastel."),
    ("Wow / high-impact", "Choose Wow Satellite Glow, especially for a dark hero with glowing map and data visuals."),
    ("Professional / academic", "Choose Professional Terra Blue as the safest default for academic credibility."),
    ("Premium / designer", "Choose Premium Editorial Contrast for a magazine-like research profile."),
    ("Recommended setup", "Default to Professional Terra Blue for credibility, then offer Dreamy Aqua Bloom as an alternate theme or softer variant."),
], header=("Client mood", "Best color direction"))

doc.add_heading("3. Improved Color Scheme Samples", level=1)
add_para(doc, "These schemes are stronger than basic blue/green academic colors. They are designed to look modern, polished, and suitable for a research portfolio.")
for palette in PALETTES:
    add_palette(doc, palette)

doc.add_heading("4. How the Colors Can Work on the Website", level=1)
add_two_col_table(doc, [
    ("Hero section", "Use the most emotional color combination here: dreamy gradient, map texture, profile photo, and clear research positioning."),
    ("Research tabs", "Each research theme can have its own accent: groundwater, AI, exposure, policy, cryosphere."),
    ("Project cards", "Use color chips for project status, contaminant type, methods, and study region."),
    ("Graphics", "Use water lines, contour maps, satellite grid overlays, and soft glow highlights instead of random abstract decoration."),
    ("Dark/light modes", "Keep one polished light theme and optionally one dramatic dark theme. Do not create too many themes unless budget allows."),
], header=("Website area", "Color usage"))

doc.add_heading("5. Style Direction Options", level=1)
add_two_col_table(doc, [
    ("Dreamy Scientific", "Soft gradients, glassy map cards, calm aqua/lavender colors, beautiful but still research-led."),
    ("Wow Research Atlas", "Dark hero, glowing satellite/map visuals, animated metrics, strong first impression."),
    ("Professional Material UI", "Clean tabs, cards, filters, publication lists, and admin-friendly editing. Best for a Next.js/CMS system."),
    ("Premium Editorial", "Large visuals, confident typography, strong contrast, and a polished publication-style layout."),
    ("Neo-Brutal Accent", "Use only in small doses: bold labels, strong borders, and bright section markers. Too much can reduce academic seriousness."),
], widths=(2.0, 4.3), header=("Style", "How to explain it"))

doc.add_heading("6. Next.js Rebuild With Editable Color/Content Options", level=1)
add_para(doc, "The new portfolio can be built in Next.js with structured content so Kavya can add or edit tabs, projects, publications, research notes, and section text without touching raw HTML.")
add_two_col_table(doc, [
    ("Editable tabs", "Research themes, project categories, publications, gallery, CV sections, and news can be schema-driven."),
    ("Color/theme control", "A simple theme config can store approved palettes. The client can choose one default theme, and a future admin option can switch between approved themes."),
    ("Content editing", "Low-cost option: JSON/MDX files. Better client option: Sanity, Contentful, or another CMS. Premium option: custom admin dashboard."),
    ("Design system", "Tailwind CSS or Material UI with design tokens for colors, spacing, cards, tabs, and buttons."),
    ("Deployment", "Vercel with SEO metadata, image optimization, responsive design, analytics, and preview links."),
], header=("Feature", "Recommendation"))

doc.add_heading("7. Partner Talking Points", level=1)
add_bullets(doc, [
    "We will not simply add random bright colors. We will create a visual identity from her research: water, maps, satellite data, AI, and environmental risk.",
    "She can choose the mood she wants: dreamy, wow, professional, soft pastel, or premium editorial.",
    "The safest recommendation is Professional Terra Blue as the default, with Dreamy Aqua Bloom as a softer alternate.",
    "The Next.js rebuild can make tabs and content editable, so future updates do not require rebuilding static HTML pages.",
    "If budget allows, we can add an interactive research atlas or dark-mode visual hero as a premium upgrade.",
])

doc.add_heading("8. Questions to Ask Kavya", level=1)
add_numbered(doc, [
    "Which mood does she prefer: dreamy, wow, professional, soft pastel, or premium editorial?",
    "Does she want one fixed color scheme or a light/dark theme option?",
    "Should she be able to edit only content, or also switch approved color themes?",
    "Which sections must be editable: research tabs, projects, publications, research notes, CV, gallery, or all of them?",
    "Will she provide updated photos, CV, publication list, and final project images?",
])

add_callout(doc, "Best final recommendation", "Pitch the project as a Next.js academic portfolio redesign with selectable color direction, strong research graphics, editable tabs, and a professional default theme. This gives her the vibrant look she wants while keeping the site credible for academic and research opportunities.", "FFF7ED", "F59E0B")

doc.save(OUT)
print(OUT)
